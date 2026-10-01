"""Deterministic verification-only PDF and editable OMML Word release."""
from pathlib import Path
import csv, re, html, os, subprocess, zipfile, io
from html.parser import HTMLParser
from datetime import datetime, timezone
import build_manuscript as pdf
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'output/release'; OUT.mkdir(parents=True,exist_ok=True)
MATH=[
r'g_0=H_0-h_0,\quad \gamma=\frac{g_0}{h_0}',
r'p(T)=p_a+\frac{(p_b-p_a)(T-T_a)}{T_b-T_a}',
r'\varepsilon=\varepsilon^e+\varepsilon^{ve}+\varepsilon^{th},\quad\varepsilon^m=\varepsilon-\varepsilon^{th}',
r'\varepsilon^{th}(T)=\alpha(T-T_r)I,\quad\alpha=68.0\times10^{-6}\ \mathrm{K}^{-1}',
r'C(E,\nu):A=2GA+\lambda\operatorname{tr}(A)I;\quad G=\frac{E}{2(1+\nu)},\quad\lambda=\frac{E\nu}{(1+\nu)(1-2\nu)}',
r'\sigma=C_\infty:\varepsilon^m+\sum_{i=1}^{23}s_i;\quad\frac{\partial s_i}{\partial\xi}+\frac{s_i}{\tau_i}=C_i:\frac{\partial\varepsilon^m}{\partial\xi}',
r'E_{inst}=E_\infty+\sum_{i=1}^{23}k_i,\quad g_i=\frac{k_i}{E_{inst}};\quad G_{inst}=\frac{E_{inst}}{2(1+\nu)},\quad K_{inst}=\frac{E_{inst}}{3(1-2\nu)}',
r'E_r(t,T)=E_\infty+\sum_{i=1}^{23}k_i\exp\left[-\frac{t}{a_T\tau_i}\right]',
r'\log_{10}a_T=\begin{cases}C_3\left(\frac{1}{T_K}-\frac{1}{T_{g,K}}\right),&T<T_g\\-\frac{C_1(T-T_g)}{C_2+(T-T_g)},&T\geq T_g\end{cases}',
r'\xi(t)=\int_0^t [a_T(T(s))]^{-1}\,ds',
r'r_i=\frac{\Delta\xi}{\tau_i},\quad A_i=\exp(-r_i),\quad B_i=\frac{1-\exp(-r_i)}{r_i};\quad s_i^{n+1}=A_i s_i^n+B_i C_i:\Delta\varepsilon^m',
r'\rho c_p\frac{\partial T}{\partial t}=\nabla\cdot(k\nabla T)',
r'\nabla\cdot\sigma=0',
r'g_n\geq0,\quad p_n\geq0,\quad g_n p_n=0',
r'\delta_q=\frac{|q_{fine}-q_{medium}|}{|q_{fine}|}\times100\%',
r'S_c(t)=1-\sum_{n=0}^{\infty}A_n\exp\left(-\frac{\zeta_n^2\alpha t}{L^2}\right),\quad\zeta_n\tan\zeta_n=Bi\\ A_n=\frac{4\sin\zeta_n}{2\zeta_n+\sin2\zeta_n},\quad T_c(t)=T_i+\int_0^t S_c(t-\tau)\,dT_\infty(\tau)',
r'u_x=Le,\ \sigma_x=0\ (\mathrm{free});\quad u_x=0,\ \sigma_x=-Ee\ (\mathrm{fixed})\\ R_x=-A\sigma_x,\quad\varepsilon_y=\varepsilon_z=e-\frac{\nu\sigma_x}{E}',
r'\sigma(t)=E_\infty v(q-a)+\sum_{i=1}^{23}E_i v\tau_i\left[\exp\left(-\frac{t-q}{\tau_i}\right)-\exp\left(-\frac{t-a}{\tau_i}\right)\right]',
r'e_i=p_i-y_i,\quad a_i=|e_i|,\quad r_i=\frac{100a_i}{|y_i|}\quad(y_i\ne0)',
r'\mathrm{MAE}=\frac{1}{n}\sum_i a_i,\quad\mathrm{RMSE}=\sqrt{\frac{1}{n}\sum_i e_i^2}',
r'a_i\leq b_{num,i}+b_{meas,i}'
]
def footer(c,doc):
    c.saveState();c.setFont('Sans',8);c.setFillColor(colors.HexColor('#52626b'))
    c.drawString(54,27,'CONSTRAINED ANNEALING  |  VERIFICATION ONLY');c.drawRightString(A4[0]-54,27,str(doc.page));c.restoreState()
def readcsv(p):
    with (ROOT/p).open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def plain(s):return html.unescape(re.sub('<[^>]+>','',s))
def supplemental(source):
    appendix=source.split('### Supplementary information',1)[1]
    text='# Supplementary information for constrained annealing verification\n\nVerification-only release · 1 October 2026\n\nThis supplement documents genuine reference calculations and published material data. No production simulation, independent validation prediction, sensitivity index, uncertainty interval, Pareto solution or optimization confirmation exists. Missing results are not zero responses. Reference numbers refer to the accompanying main article.\n\n'+appendix
    text+='\n### S3. Full scalar verification and refinement records\n\nThe following tables reproduce every accepted scalar comparison and each audited mesh/time response record. Raw nodal and element exports, complete input decks, warnings and rejected attempts remain in the repository at the paths indexed by result_traceability.csv. Units follow each quantity; reaction is force in the structural verification and force per unit depth in plane-strain contact convergence. Percent errors at a zero reference are undefined, not zero.\n'
    table_specs=[('verification/structural_verification.csv','All structural reference comparisons',['case','time_s','quantity','ansys','reference','absolute_error'],['Case','t (s)','Quantity','ANSYS','Reference','Absolute error']),('verification/contact_verification.csv','All gap-contact reference comparisons',['case','time_s','quantity','ansys','reference','absolute_error'],['Case','t (s)','Quantity','ANSYS','Reference','Absolute error']),('convergence/stage_19_audited_refinement.csv','All audited mesh and time refinement responses',['table','run_path','quantity','value','unit','delta_percent'],['Study','Case / level','Quantity','Value','Unit','Change (%)'])]
    def nice(v):
        try:
            n=float(v)
            if n==0:return '0'
            s=f'{n:.7g}'
            if 'e' in s:
                mant,exp=s.split('e');return mant+' × 10<sup>'+str(int(exp)).replace('-','−')+'</sup>'
            return s.replace('-','−')
        except (TypeError,ValueError):return str(v)
    def label(v):
        v=v.replace('degC/m','K m⁻¹').replace('degC','°C').replace('N/mm','N mm⁻¹')
        for a,b in [('mesh_convergence.csv','Mesh'),('timestep_convergence.csv','Time'),('simulation/convergence/stage10_',''),('structural_','S '),('thermal_','T '),('contact_','C '),('mesh_',''),('time_',''),('extra_fine','extra fine'),('ultra_fine','ultra fine'),('eigen_free','eigen free'),('eigen_fixed','eigen fixed'),('W_max','W<sub>max</sub>'),('residual','resid.'),('maximum','max.'),('temperature','temp.'),('gradient','grad.'),('pressure','press.'),('displacement','disp.'),('stress','stress'),('_MPa',' (MPa)'),('_mm',' (mm)'),('_N',' (N)'),('_',' ')]:v=v.replace(a,b)
        return v
    for path,title,keys,heads in table_specs:
        text+='\nSUPPTABLE: '+title+'\n'+' | '.join(heads)+'\n'
        for row in readcsv(path):text+=' | '.join(label(nice(row[k])) for k in keys)+'\n'
    text+='\n### S4. Material property inventory and use restrictions\n\nThe property registry below preserves all recorded numerical PLA values, including comparator formulations that are not admitted to the reference material card. P = Prusament PLA; all other formulations are comparator evidence. Exact print settings, confidence, full source locators, temperature context and ANSYS-use restrictions are retained in material/pla_properties.csv. A number in this inventory is not authorization to mix formulations. Tables S2–S3 provide the reference spectrum and candidate fixture data.\n'
    text+='\nSUPPTABLE: Complete PLA property inventory with source and formulation\nID / source | Property / component | Value / unit | Temperature / formulation\n'
    for r in readcsv('material/pla_properties.csv'):
        sym=r['property_symbol'];component=r['component']
        for token in ['sym','component']:
            val=sym if token=='sym' else component
            val=re.sub(r'_([A-Za-z0-9∞]+)',r'<sub>\1</sub>',val)
            val=re.sub(r'\^([A-Za-z0-9]+)',r'<sup>\1</sup>',val)
            if token=='sym':sym=val
            else:component=val
        temp=r['temperature_C']+('' if r['temperature_C']=='not reported' else ' °C')
        unit='' if r['unit']=='1' else ' '+r['unit']
        text+=' | '.join([r['record_id']+'<br/>'+r['source_key'],sym+' / '+component,nice(r['value'])+unit,temp+'<br/>'+r['PLA_formulation']])+'\n'
    text+='\n### S5. Geometry, discretization and run matrix\n\nThe intended coupon is 60 mm × 10 mm × 4 mm, with two 70 mm × 20 mm × 5 mm plates. The reference-state candidate total clearances are 0, 0.01, 0.02, 0.04 and 0.08 mm. These are design choices. No production geometry mesh, heat-transfer boundary history or released-state protocol is approved.\n\nThe complete accepted verification matrix comprises one plane-wall thermal run, six structural/contact cases and 28 unique convergence runs. The 43 convergence attempt directories include 15 excluded attempts; they are not independent experiments. The table below identifies every accepted convergence run. Complete case parameters, mesh counts, time histories, contact controls and solver messages are preserved in each case.json, input.dat and mapdl.out.\n'
    runs=sorted({r['run_path'] for r in readcsv('convergence/mesh_convergence.csv')+readcsv('convergence/timestep_convergence.csv')+readcsv('convergence/contact_sensitivity.csv')})
    text+='\nSUPPTABLE: Complete accepted convergence case matrix\nRun directory beneath simulation/convergence | Scope\n'
    for run in runs:text+=run.split('/')[-1].replace('_',' ')+' | Reference discretization / contact study; no production prediction\n'
    text+='\n### S6. Validation and unavailable analyses\n\nThe validation registry contains 45 literature observations: 18 reserved source-specific means, 24 excluded context observations and three quarantined secondary maxima. No prediction or validation error is filled. The six temperature levels in Table S4 describe a separate published Ultimaker experiment and are not the proposed Prusament production matrix. Source material, cooling and thermal-contact closure remain insufficient for an independent solve.\n\nThe production design and case manifest, all-cases results, Pareto and confirmation tables contain zero result rows. There are no surrogate diagnostics, global sensitivity indices, material probability distributions, uncertainty intervals, Pareto solutions or confirmation runs. No additional production contours can be provided. This release does not claim completion of these studies.\n\n### S7. Reproduction and interpretation\n\nUse Python 3.12 with the pinned requirements. The release was built with Python 3.12.14, ReportLab 4.4.9, python-docx 1.2.0, lxml 6.1.1 and latex2mathml 3.78.1. Windows fonts and the installed Office MathML-to-OMML stylesheet are required for the present export script; these are not redistributed. Word equations are native editable OMML.\n\nRun the unit tests, check_release.py, and the commands in reproducibility/README.md. Archived solver decks require compatible Ansys Student 2026 R1 / MAPDL 26.1 update 20260202 and a usable license. A fresh result directory is mandatory. Do not overwrite archived attempts. Rerun comparison calculations before accepting newly solved cases. A later solver version need not produce identical binary files.\n\nThe SHA-256 manifest covers the final three files, critical data, input decks, analysis and build scripts. The traceability and figure/table manifests identify evidence paths. Historical third-party PDFs have been removed from the current release tree and remain unchanged locally; earlier Git history was not rewritten. Obtain those sources legally using external_evidence.csv. Published values are not new experiments. No scientific image was generated by AI.\n'
    return text

class Inline(HTMLParser):
    def __init__(self,p):super().__init__();self.p=p;self.tags=[]
    def handle_starttag(self,t,a):
        if t=='br':self.p.add_run().add_break()
        else:self.tags.append(t)
    def handle_endtag(self,t):
        if t in self.tags:self.tags.remove(t)
    def handle_data(self,s):
        for part in re.split(r'([₀₁₂₃₄₅₆₇₈₉ᵢ∞]+|[⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺]+)',s):
            if not part:continue
            r=self.p.add_run(part)
            if 'b' in self.tags:r.bold=True
            if 'i' in self.tags:r.italic=True
            r.font.subscript='sub' in self.tags;r.font.superscript='sup' in self.tags or 'super' in self.tags
            if all(c in '₀₁₂₃₄₅₆₇₈₉ᵢ' for c in part):r.text=part.translate(str.maketrans('₀₁₂₃₄₅₆₇₈₉ᵢ','0123456789i'));r.font.subscript=True
            if all(c in '⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺' for c in part):r.text=part.translate(str.maketrans('⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺','0123456789−+'));r.font.superscript=True
def addtext(p,s):Inline(p).feed(s)
def bookmark(p,name,num):
    # Number-only bookmarks make REF fields stay concise after Word updates them.
    runs=p.runs
    target=runs[-1] if name.startswith('Equation') else runs[0]
    if name.startswith(('Figure','Table')):
        label=re.match(r'(Figure \d+|Table \d+)',target.text).group(1)
        tail=target.text[len(label):];target.text=label
        extra=p.add_run(tail);target._r.addnext(extra._r)
    a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(num));a.set(qn('w:name'),name);target._r.addprevious(a)
    b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(num));target._r.addnext(b)
def field(p,instruction,cached):
    node=OxmlElement('w:fldSimple');node.set(qn('w:instr'),instruction)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=cached;r.append(t);node.append(r);p._p.append(node)
def docx(source,geometry):
    d=Document();sec=d.sections[0];sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
    sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(11)
        st.paragraph_format.space_after=Pt(6);st.paragraph_format.widow_control=True
        for borders in list(st.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
        for fonts in st.element.iter(qn('w:rFonts')):
            for key in list(fonts.attrib):
                if key.endswith('Theme'):del fonts.attrib[key]
    d.styles['Normal'].paragraph_format.line_spacing=1.05
    for name,size in [('Title',20),('Heading 1',13),('Heading 2',11.5)]:
        d.styles[name].font.size=Pt(size);d.styles[name].font.bold=True;d.styles[name].paragraph_format.keep_with_next=True
    d.styles['Caption'].font.size=Pt(9)
    p=sec.footer.paragraphs[0];p.add_run('CONSTRAINED ANNEALING | VERIFICATION ONLY     ').font.size=Pt(8);field(p,'PAGE','1')
    stylesheet=Path(os.environ.get('MML2OMML_XSL',r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
    transform=etree.XSLT(etree.parse(str(stylesheet)))
    lines=source.splitlines();i=0;eqn=0;tabn=0
    while i<len(lines):
        s=lines[i].strip();i+=1
        if not s:continue
        if s=='<!-- PAGE -->':d.add_page_break();continue
        if s.startswith('EQ:'):
            eqn+=1;p=d.add_paragraph();p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(7)
            # Use separate native display lines where the journal equation has two rows.
            chunks=MATH[eqn-1].split('\\\\') if eqn in [16,17] else [MATH[eqn-1]]
            for j,tex in enumerate(chunks):
                if j:p.add_run().add_break()
                math=etree.fromstring(convert(tex).encode());omml=transform(math).getroot()
                # Display math with explicit Cambria Math and a size that fits long relations.
                for mr in omml.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}r'):
                    rp=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');rp.append(fonts)
                    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20' if eqn in [5,7,11,18] else '22');rp.append(sz);mr.insert(0,rp)
                p._p.append(omml)
            p.add_run(f'   ({eqn})');bookmark(p,'Equation'+str(eqn),300+eqn);continue
        if s.startswith('TABLE:'):
            tabn+=1;p=d.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;addtext(p,f'Table {tabn}. '+s[6:].strip());bookmark(p,'Table'+str(tabn),100+tabn)
            rr=[]
            while i<len(lines) and ' | ' in lines[i]:rr.append(lines[i].split(' | '));i+=1
            t=d.add_table(rows=0,cols=len(rr[0]));t.autofit=False
            widths=[1.65,2.0,3.1] if len(rr[0])==3 else [1.15,1.65,1.75,2.2]
            borders=OxmlElement('w:tblBorders')
            for tag in ['top','left','bottom','right','insideH','insideV']:
                e=OxmlElement('w:'+tag);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            t._tbl.tblPr.append(borders)
            for ri,row in enumerate(rr):
                cells=t.add_row().cells
                no_split=OxmlElement('w:cantSplit');t.rows[-1]._tr.get_or_add_trPr().append(no_split)
                if ri==0:rep=OxmlElement('w:tblHeader');t.rows[-1]._tr.get_or_add_trPr().append(rep)
                for ci,v in enumerate(row):
                    cell=cells[ci];cell.width=Inches(widths[ci]);addtext(cell.paragraphs[0],v)
                    if ri==0:cell.paragraphs[0].paragraph_format.keep_with_next=True
                    margins=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
                    cell._tc.get_or_add_tcPr().append(margins)
                    for r in cell.paragraphs[0].runs:r.font.size=Pt(9);r.bold=ri==0
                    if ri==0:fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E9EFF2');cell._tc.get_or_add_tcPr().append(fill)
            d.add_paragraph();continue
        if s.startswith(('FIGURE:','IMAGE:')):
            path=geometry if s.startswith('FIGURE:') else ROOT/s[6:].strip()
            p=d.add_paragraph();p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Inches(6.55));continue
        if s.startswith('CAPTION:'):
            p=d.add_paragraph(style='Caption');addtext(p,s[8:].strip());n=int(re.search(r'Figure (\d+)',s).group(1));bookmark(p,'Figure'+str(n),200+n);continue
        if s.startswith('# '):p=d.add_paragraph(style='Title');addtext(p,s[2:]);continue
        if s.startswith('### '):
            h=s[4:];style='Heading 2' if re.match(r'\d+\.\d+\.',h) else 'Heading 1';p=d.add_paragraph(style=style);addtext(p,h);continue
        p=d.add_paragraph();addtext(p,s)
        if re.match(r'^\[\d+\]',s):bookmark(p,'Reference'+re.match(r'^\[(\d+)\]',s).group(1),400+int(re.match(r'^\[(\d+)\]',s).group(1)));p.paragraph_format.keep_together=True
    # An editable internal index supplies actual REF fields, not simulated field text.
    d.add_heading('Cross reference index',level=1)
    p=d.add_paragraph('Figures: ')
    for n in range(1,5):field(p,f'REF Figure{n} \\h',f'Figure {n}');p.add_run('; ')
    p=d.add_paragraph('Tables: ')
    for n in range(1,4):field(p,f'REF Table{n} \\h',f'Table {n}');p.add_run('; ')
    p=d.add_paragraph('Equations: ')
    for n in range(1,22):field(p,f'REF Equation{n} \\h',f'({n})');p.add_run(' ')
    d.core_properties.title='Thermal viscoelastic and contact verification toward gap constrained annealing of FFF printed PLA'
    d.core_properties.subject='Verification-only package; production study incomplete'
    d.core_properties.author='Aadit Jain; Dheeraj Yadav; Ekansh Malhotra';d.core_properties.created=d.core_properties.modified=datetime(2026,10,1,tzinfo=timezone.utc)
    for borders in list(d.element.iter(qn('w:pBdr'))):borders.getparent().remove(borders)
    dest=OUT/'Constrained-Annealing-2026.docx';d.save(dest)
    # Canonical ZIP timestamps make repeated exports byte-identical.
    original=dest.read_bytes();buf=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as a,zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as b:
        for name in sorted(a.namelist()):
            info=zipfile.ZipInfo(name,(2026,10,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;b.writestr(info,a.read(name))
    dest.write_bytes(buf.getvalue());print(dest)
def main():
    source=(ROOT/'manuscript/current.md').read_text(encoding='utf-8')
    main=source.split('### Supplementary information',1)[0].rstrip()
    main=re.sub(r'<!-- PAGE -->\s*$','',main).rstrip()
    supplement='\n'.join(line.rstrip() for line in supplemental(source).splitlines())+'\n';(ROOT/'supplementary').mkdir(exist_ok=True)
    (ROOT/'supplementary/current.md').write_text(supplement,encoding='utf-8',newline='\n')
    pdf.footer=footer
    for name in ['title','subtitle','heading']:pdf.ST[name].textColor=colors.black
    pdf.OUT=OUT/'Constrained-Annealing-2026.pdf';pdf.build(main)
    pdf.CAPTION_MIN_SPACE=110;pdf.ST['tablecaption'].keepWithNext=False
    pdf.OUT=OUT/'Constrained-Annealing-2026-Supplementary.pdf';pdf.build(supplement)
    pdf.CAPTION_MIN_SPACE=0;pdf.ST['tablecaption'].keepWithNext=True
    folder=ROOT/'figures/release';folder.mkdir(exist_ok=True);vector=folder/'geometry.pdf'
    c=Canvas(str(vector),pagesize=(487,245),invariant=1);g=pdf.GapFigure();g.drawOn(c,6,8);c.save()
    poppler=os.environ.get('PDFTOPPM',r'C:\Users\asus\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdftoppm.exe')
    subprocess.run([poppler,'-png','-singlefile','-r','240',str(vector),str(folder/'geometry')],check=True,capture_output=True)
    docx(main,folder/'geometry.png')
    pdf.OUT=ROOT/'output/pdf/Constrained-Annealing-2026-DRAFT.pdf';pdf.build(source)
if __name__=='__main__':main()
